"""
Energy report generation — real files, not placeholders.

Assembles data from energy_registry / energy_history / energy_optimizer
into a report-type-specific "sections" structure, then renders it to bytes
in either PDF (reportlab) or Excel (openpyxl). Both renderers consume the
exact same sections structure, so adding a new report type only means
adding one branch to _gather_report_data() — the rendering code never
needs to change.
"""

from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import energy_history, energy_optimizer, energy_registry

REPORT_TITLES = {
    "daily": "Daily Energy Report",
    "weekly": "Weekly Energy Report",
    "monthly": "Monthly Energy Report",
    "cost": "Cost Analysis Report",
    "efficiency": "Energy Efficiency Report",
}


def _today_str() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def _device_row(device_id: str) -> str:
    return device_id.replace("_", " ").title()


def _gather_report_data(report_type: str) -> dict:
    now = datetime.now()
    summary = energy_history.get_summary()
    sections = []

    if report_type == "daily":
        live = energy_registry.get_live()
        devices = energy_registry.get_devices()
        sections.append(
            {
                "heading": "Today's Snapshot",
                "kv": [
                    ("Voltage", f"{live['voltage']} V"),
                    ("Power", f"{live['power']} W"),
                    ("Energy (session)", f"{live['energy_kwh']} kWh"),
                    ("Power Factor", live["power_factor"]),
                    ("System Status", live["system_status"]),
                ],
            }
        )
        sections.append(
            {
                "heading": "Device Breakdown",
                "table_headers": ["Device", "Power (W)", "Daily Energy (kWh)", "Status"],
                "table_rows": [
                    [_device_row(d["id"]), d["power_w"], d["daily_energy_kwh"], "ON" if d["is_on"] else "OFF"]
                    for d in devices
                ],
            }
        )

    elif report_type in ("weekly", "monthly"):
        days = 7 if report_type == "weekly" else 30
        start = (now - timedelta(days=days - 1)).strftime("%Y-%m-%d")
        end = _today_str()
        series = energy_history.get_series("daily", start, end)
        sections.append(
            {
                "heading": f"Consumption — last {days} days",
                "table_headers": ["Date", "Energy (kWh)", "Cost"],
                "table_rows": [[s["date"], s["total_kwh"], f"${s['cost']}"] for s in series],
            }
        )
        comparison = energy_history.get_device_comparison(start, end)
        sections.append(
            {
                "heading": "Device Comparison",
                "table_headers": ["Device", "Total Energy (kWh)"],
                "table_rows": [[_device_row(c["device_id"]), c["total_kwh"]] for c in comparison],
            }
        )

    elif report_type == "cost":
        start = (now - timedelta(days=29)).strftime("%Y-%m-%d")
        end = _today_str()
        comparison = energy_history.get_device_comparison(start, end)
        sections.append(
            {
                "heading": "Cost Breakdown by Device (last 30 days)",
                "table_headers": ["Device", "Energy (kWh)", "Estimated Cost"],
                "table_rows": [
                    [
                        _device_row(c["device_id"]),
                        c["total_kwh"],
                        f"${round(c['total_kwh'] * energy_optimizer.COST_PER_KWH, 2)}",
                    ]
                    for c in comparison
                ],
            }
        )
        sections.append(
            {
                "heading": "Summary",
                "kv": [
                    ("Estimated Cost Saved (year)", f"${summary['estimated_cost_saved']}"),
                    ("Rate Used", f"${energy_optimizer.COST_PER_KWH}/kWh"),
                ],
            }
        )

    elif report_type == "efficiency":
        recs = energy_optimizer.generate_recommendations()
        sections.append(
            {
                "heading": "Efficiency Summary",
                "kv": [
                    ("Efficiency Score", f"{summary['efficiency_score_pct']}%"),
                    ("Energy Saved (month)", f"{summary['monthly_energy_saved_kwh']} kWh"),
                    ("CO2 Reduced (year)", f"{summary['estimated_co2_reduced_kg']} kg"),
                ],
            }
        )
        sections.append(
            {
                "heading": "AI Recommendations",
                "table_headers": ["Recommendation", "Expected Saving", "Confidence"],
                "table_rows": [
                    [r["title"], f"{r['expected_saving_pct']}%", f"{int(r['confidence_score'] * 100)}%"]
                    for r in recs
                ]
                or [["No active recommendations", "-", "-"]],
            }
        )

    return {
        "title": REPORT_TITLES.get(report_type, "Energy Report"),
        "generated_at": now.strftime("%Y-%m-%d %H:%M"),
        "sections": sections,
    }


def build_pdf(report_type: str) -> bytes:
    data = _gather_report_data(report_type)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"HydroMind AI — {data['title']}", styles["Title"]),
        Paragraph(f"Generated: {data['generated_at']}", styles["Normal"]),
        Spacer(1, 0.5 * cm),
    ]

    for section in data["sections"]:
        story.append(Paragraph(section["heading"], styles["Heading2"]))
        if "kv" in section:
            rows = [[k, str(v)] for k, v in section["kv"]]
            table = Table(rows, colWidths=[8 * cm, 8 * cm])
            table.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9)]))
        else:
            rows = [section["table_headers"]] + [[str(c) for c in row] for row in section["table_rows"]]
            table = Table(rows, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2a78d6")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e1e0d9")),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9f9f7")]),
                    ]
                )
            )
        story.append(table)
        story.append(Spacer(1, 0.5 * cm))

    doc.build(story)
    return buffer.getvalue()


def build_xlsx(report_type: str) -> bytes:
    data = _gather_report_data(report_type)
    wb = Workbook()
    ws = wb.active
    ws.title = data["title"][:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2A78D6", end_color="2A78D6", fill_type="solid")

    ws.append([f"HydroMind AI — {data['title']}"])
    ws.append([f"Generated: {data['generated_at']}"])
    ws.append([])

    for section in data["sections"]:
        ws.append([section["heading"]])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)

        if "kv" in section:
            for k, v in section["kv"]:
                ws.append([k, v])
        else:
            ws.append(section["table_headers"])
            header_row = ws.max_row
            for c in range(1, len(section["table_headers"]) + 1):
                cell = ws.cell(row=header_row, column=c)
                cell.font = header_font
                cell.fill = header_fill
            for row in section["table_rows"]:
                ws.append(row)
        ws.append([])

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(40, length + 4)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
